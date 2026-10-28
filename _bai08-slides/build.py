"""Dựng lại Bài 8 – Hệ thống gợi ý (bản giảng viên + bản sinh viên).

Chạy: python3 -I build.py <thư mục media gốc> <thư mục output>
"""
import sys, copy
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

MEDIA, OUT = sys.argv[1], sys.argv[2]

NAVY = "1C3A6E"; YEL = "F5C518"; YEL_T = "FFF6D6"; TINT = "EAF0F8"
INK = "26303D"; MUTED = "5F6B7A"; RED = "B03A5B"; GREEN = "2E7D57"; LINE = "C9D3E1"
WHITE = "FFFFFF"
FONT = "Arial"
SW, SH = 12188825, 6858000
L = Inches(0.5); CW = SW - 2 * L  # lề trái và bề rộng nội dung

IMG = {
    "cover": "image1.png", "white": "image2.png", "blue": "image4.png",
    "man_hinh": "image5.png", "phan_bo": "image6.png", "ma_tran": "image9.png",
    "duoi_dai": "image10.png", "thanh_phan": "image11.png", "ma_tran_8x8": "image13.png",
    "ho_so_an": "image14.png", "do_giong": "image15.png", "lang_gieng": "image16.png",
    "cum_phim": "image17.png", "prize": "image18.png", "ban_do": "image19.png",
    "so_sanh": "image21.png", "vong": "image24.png",
}


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, teacher):
        self.teacher = teacher
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]

    # ---- khung ----
    def slide(self, bg="white"):
        s = self.prs.slides.add_slide(self.blank)
        s.shapes.add_picture(f"{MEDIA}/{IMG[bg]}", 0, 0, SW, SH)
        return s

    def notes(self, s, text):
        if self.teacher and text:
            s.notes_slide.notes_text_frame.text = text

    # ---- chữ ----
    def text(self, s, x, y, w, h, paras, size=16, color=INK, bold=False, align="l",
             anchor="t", italic=False, space_after=0, line=1.1, bullet=False, margin=0):
        """paras: str | list[str | list[(text, opts)]]"""
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, margin)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str):
            paras = [paras]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            para.line_spacing = line
            para.space_after = Pt(space_after)
            if bullet:
                self._bullet(para, size)
            runs = [(p, {})] if isinstance(p, str) else p
            for t, o in runs:
                r = para.add_run()
                r.text = t
                f = r.font
                f.name = FONT
                f.size = Pt(o.get("size", size))
                f.bold = o.get("bold", bold)
                f.italic = o.get("italic", italic)
                f.color.rgb = rgb(o.get("color", color))
        return tb

    @staticmethod
    def _bullet(para, size):
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Pt(size) * 1.1)))
        pPr.set("indent", str(-Emu(Pt(size) * 1.1)))
        bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", FONT)
        ch = etree.SubElement(pPr, qn("a:buChar")); ch.set("char", "•")

    def title(self, s, t, color=NAVY):
        self.text(s, L, Inches(0.62), CW, Inches(0.7), t, size=28, bold=True, color=color, anchor="m")

    def source(self, s, t):
        """Chú thích nguồn/minh hoạ ở chân slide, phía phải dải vàng của nền."""
        self.text(s, Inches(6.4), Inches(6.9), Inches(6.43), Inches(0.3), t, size=11,
                  color=MUTED, align="r", anchor="b")

    def box(self, s, x, y, w, h, fill=TINT, line=None, radius=0.06):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        if radius:
            shp.adjustments[0] = radius
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.text_frame.text = ""
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
        return shp

    def hline(self, s, x, y, w, color=LINE, weight=0.75):
        ln = s.shapes.add_connector(1, x, y, x + w, y)
        ln.line.color.rgb = rgb(color); ln.line.width = Pt(weight)
        st = ln._element.find(qn("p:style"))
        if st is not None:
            ln._element.remove(st)
        return ln

    def question(self, s, label, t, y=Inches(5.95), h=Inches(0.95)):
        self.box(s, L, y, CW, h, fill=YEL_T, radius=0.08)
        self.text(s, L + Inches(0.3), y, CW - Inches(0.6), h,
                  [[(label + " ", {"bold": True, "color": NAVY}), (t, {})]],
                  size=18, anchor="m")

    def takeaway(self, s, t, y=Inches(6.2), h=Inches(0.7)):
        self.box(s, L, y, CW, h, fill=NAVY, radius=0.08)
        self.text(s, L + Inches(0.3), y, CW - Inches(0.6), h, t, size=18, bold=True,
                  color=WHITE, anchor="m")

    def pic(self, s, key, x, y, w=None, h=None):
        return s.shapes.add_picture(f"{MEDIA}/{IMG[key]}", x, y, w, h)

    def table(self, s, x, y, w, rows, col_w, size=15, header=True, row_h=Inches(0.55),
              first_col_bold=False, highlight=None):
        shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
        tbl = shp.table
        tblPr = tbl._tbl.tblPr
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
        for i, cwid in enumerate(col_w):
            tbl.columns[i].width = cwid
        for r, row in enumerate(rows):
            tbl.rows[r].height = min(row_h, Inches(0.6)) if header and r == 0 else row_h
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.12)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                is_head = header and r == 0
                fill = NAVY if is_head else (YEL_T if highlight and (r, c) in highlight else (WHITE if r % 2 else "F4F7FB"))
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run(); run.text = val
                f = run.font; f.name = FONT; f.size = Pt(size)
                f.bold = is_head or (first_col_bold and c == 0)
                f.color.rgb = rgb(WHITE if is_head else (NAVY if first_col_bold and c == 0 else INK))
                # chỉ kẻ đường ngang mảnh
                tcPr = cell._tc.get_or_add_tcPr()
                for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.SubElement(tcPr, qn(tag))
                    if tag in ("a:lnL", "a:lnR"):
                        ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
                    else:
                        ln.set("w", "9525"); sf = etree.SubElement(ln, qn("a:solidFill"))
                        c_ = etree.SubElement(sf, qn("a:srgbClr")); c_.set("val", LINE)
                # thứ tự phần tử trong tcPr: các ln phải đứng trước fill
                fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
                for e in fills:
                    tcPr.remove(e); tcPr.append(e)
        return shp

    def hbar(self, s, x, y, label, value, frac, maxw, color, label_w=Inches(3.0), h=Inches(0.55), vsize=24):
        self.text(s, x, y, label_w - Inches(0.2), h, label, size=16, anchor="m", align="r")
        bw = max(int(maxw * frac), Inches(0.04))
        self.box(s, x + label_w, y, bw, h, fill=color, radius=0)
        self.text(s, x + label_w + bw + Inches(0.15), y, Inches(2.5), h, value, size=vsize,
                  bold=True, color=color, anchor="m")

    def divider(self, n, t, sub):
        s = self.slide("blue")
        self.text(s, Inches(1.0), Inches(2.2), Inches(1.6), Inches(2.0), str(n), size=120,
                  bold=True, color=YEL, anchor="m")
        self.text(s, Inches(2.7), Inches(2.45), Inches(9.6), Inches(1.0), t, size=36,
                  bold=True, color=WHITE, anchor="b")
        self.text(s, Inches(2.7), Inches(3.55), Inches(9.6), Inches(0.8), sub, size=18,
                  color="C9D6EA", anchor="t")
        return s


QUESTIONS = [
    ("Hệ gợi ý giải bài toán gì?", "Trang chủ Netflix: vài chục ô, hàng nghìn phim"),
    ("Hệ gợi ý học từ dữ liệu nào?", "Từ 5 sao sang Thích / Không thích"),
    ("Gợi ý dựa vào nội dung hay vào người khác?", "Hàng “Vì bạn đã xem…”, dữ liệu MovieLens"),
    ("Làm sao biết hệ gợi ý tốt?", "Netflix Prize, thử lại trên MovieLens"),
    ("Sản phẩm mới, người dùng mới thì sao?", "Phim mới ra mắt, chưa ai xem"),
    ("Giá trị kinh doanh đo bằng gì?", "A/B test và tỷ lệ giữ chân của Netflix"),
]


def build(teacher):
    d = Deck(teacher)
    T = d.text

    # 1 ── Bìa
    s = d.slide("cover")
    T(s, Inches(1.0), Inches(1.75), Inches(11.3), Inches(0.5), "Phân tích dữ liệu và ra quyết định · AA6060",
      size=18, color="C9D6EA", align="c")
    T(s, Inches(1.0), Inches(2.3), Inches(11.3), Inches(1.0), "Bài 8. Hệ thống gợi ý", size=44,
      bold=True, color=WHITE, align="c")
    T(s, Inches(1.0), Inches(3.35), Inches(11.3), Inches(0.6),
      "Lọc theo nội dung, lọc cộng tác và cách đánh giá", size=22, color=YEL, align="c")
    T(s, Inches(1.0), Inches(4.3), Inches(11.3), Inches(0.5),
      "Ví dụ minh hoạ: Netflix và bộ dữ liệu MovieLens 1M", size=16, color="C9D6EA", align="c")
    T(s, Inches(1.0), Inches(5.3), Inches(11.3), Inches(0.5), "Giảng viên: Trương Hoàng Giang",
      size=16, color=WHITE, align="c")
    d.notes(s, "Gửi trước cho lớp: học liệu Bài 8 và bài đọc AA6060_Bai08_Case-study_Netflix.pdf. "
               "Bài học đi theo sáu câu hỏi về hệ gợi ý; Netflix và MovieLens chỉ là ví dụ để trả lời từng câu. "
               "Các hình màn hình ứng dụng là hình tự vẽ để minh hoạ, không phải ảnh chụp Netflix.")

    # 2 ── Câu hỏi mở đầu
    s = d.slide()
    d.title(s, "Tối nay xem phim gì?")
    T(s, L, Inches(1.5), Inches(11.0), Inches(1.0),
      "Bạn chỉ được nghe MỘT người trước khi chọn phim. Bạn nghe ai?", size=22)
    for i, (k, h, sub) in enumerate([("A", "Nhà phê bình điện ảnh", "xem rất nhiều phim, viết bài chuyên môn"),
                                     ("B", "Một người có gu giống bạn", "đã thích đúng những phim bạn từng thích")]):
        x = L + i * (CW // 2 + Inches(0.15))
        w = CW // 2 - Inches(0.15)
        d.box(s, x, Inches(2.65), w, Inches(2.6), fill=TINT)
        T(s, x + Inches(0.4), Inches(2.9), Inches(1.2), Inches(1.2), k, size=60, bold=True, color=NAVY)
        T(s, x + Inches(0.4), Inches(4.05), w - Inches(0.8), Inches(0.5), h, size=22, bold=True, color=NAVY)
        T(s, x + Inches(0.4), Inches(4.55), w - Inches(0.8), Inches(0.5), sub, size=16, color=MUTED)
    d.notes(s, "Cho lớp giơ tay A hoặc B, ghi số lên bảng. Hỏi 1–2 bạn vì sao. Chưa chốt.")

    # 3 ── Hill 1995
    s = d.slide()
    d.title(s, "Người có gu giống bạn đoán đúng hơn nhà phê bình")
    T(s, L, Inches(1.45), Inches(11.5), Inches(0.9),
      "Hill và cộng sự (1995): 291 người, 1.750 phim, 55.000 lượt chấm. Dự đoán điểm từng người sẽ chấm, "
      "rồi so với điểm họ chấm thật sau đó.", size=18, color=INK)
    T(s, L, Inches(2.55), Inches(8), Inches(0.4), "Tương quan giữa dự đoán và điểm chấm thật",
      size=14, color=MUTED)
    maxw = Inches(7.0)
    d.hbar(s, L, Inches(3.1), "Từ người có gu giống", "0,62", 0.62, maxw, NAVY, h=Inches(0.7), vsize=32)
    d.hbar(s, L, Inches(4.0), "Từ nhà phê bình", "0,22", 0.22, maxw, "8C9AAD", h=Inches(0.7), vsize=32)
    d.takeaway(s, "Đây là ý tưởng gốc của lọc cộng tác, phương pháp đứng sau phần lớn hệ gợi ý hiện nay")
    d.source(s, "Hill, Stead, Rosenstein & Furnas (CHI 1995)")
    d.notes(s, "0,62 là tương quan khá mạnh, 0,22 là yếu. Chỉ nêu tên “lọc cộng tác”, chưa giải thích; phần 3 quay lại.")

    # 4 ── Lộ trình
    s = d.slide()
    d.title(s, "Sáu câu hỏi của buổi học")
    T(s, Inches(1.35), Inches(1.45), Inches(6.3), Inches(0.4), "Câu hỏi về hệ gợi ý", size=14, bold=True, color=MUTED)
    T(s, Inches(7.9), Inches(1.45), Inches(4.9), Inches(0.4), "Ví dụ dùng để trả lời", size=14, bold=True, color=MUTED)
    y0, rh = Inches(1.9), Inches(0.78)
    for i, (q, ex) in enumerate(QUESTIONS):
        y = y0 + i * rh
        d.hline(s, L, y, CW)
        T(s, L, y, Inches(0.7), rh, str(i + 1), size=26, bold=True, color=YEL if False else NAVY, anchor="m")
        T(s, Inches(1.35), y, Inches(6.4), rh, q, size=20, bold=True, color=NAVY, anchor="m")
        T(s, Inches(7.9), y, Inches(4.9), rh, ex, size=15, color=MUTED, anchor="m")
    d.hline(s, L, y0 + 6 * rh, CW)
    d.source(s, "Netflix: Gomez-Uribe & Hunt (2015), ACM TMIS — số liệu do Netflix tự công bố")
    d.notes(s, "Mỗi phần mở bằng một câu hỏi về hệ gợi ý nói chung. Netflix/MovieLens chỉ là ví dụ để thấy câu trả lời; "
               "cuối mỗi phần chốt lại khái niệm áp dụng cho mọi doanh nghiệp (thương mại điện tử, âm nhạc, tin tức…).")

    # ── PHẦN 1
    s = d.divider(1, QUESTIONS[0][0], "Gợi ý không phải là đoán sở thích, mà là chọn thứ để hiển thị")

    s = d.slide()
    d.title(s, "Màn hình có vài chục ô, kho có hàng nghìn phim")
    d.pic(s, "man_hinh", L, Inches(1.45), w=CW)
    d.question(s, "Câu hỏi 1.", "Đây là bài toán dự đoán hay bài toán ra quyết định? Nguồn lực nào đang khan hiếm?",
               y=Inches(5.95), h=Inches(0.85))
    d.source(s, "Minh hoạ: giao diện kiểu Netflix, tự vẽ")
    d.notes(s, "Đáp: bài toán ra quyết định — chọn sản phẩm nào vào ô nào cho từng người. Khan hiếm: số ô trên màn hình "
               "và 60–90 giây chú ý của người dùng. Dự đoán “người này thích gì” chỉ là đầu vào cho quyết định.")

    s = d.slide()
    d.title(s, "Hệ gợi ý chọn K trong hàng nghìn sản phẩm cho từng người")
    d.pic(s, "phan_bo", Inches(1.4), Inches(1.45), w=Inches(10.5))
    w2 = (CW - Inches(0.6)) // 3
    for i, (h, t) in enumerate([("Đầu vào", "ai đang xem, đã làm gì, sản phẩm có gì"),
                                ("Đầu ra", "danh sách ngắn, có thứ tự"),
                                ("Ràng buộc", "chỉ có K ô; thêm sản phẩm này là bớt sản phẩm khác")]):
        x2 = L + i * (w2 + Inches(0.3))
        T(s, x2, Inches(4.5), w2, Inches(0.4), h, size=18, bold=True, color=NAVY)
        T(s, x2, Inches(4.9), w2, Inches(0.8), t, size=16, color=INK)
    d.takeaway(s, "Gợi ý là một quyết định phân bổ, giống các bài toán phân bổ nguồn lực đã học", y=Inches(6.0), h=Inches(0.75))
    d.notes(s, "Định nghĩa: hệ gợi ý dùng dữ liệu để chọn ra danh sách ngắn những sản phẩm mà từng người nhiều khả năng quan tâm. "
               "Trước đây người bán hàng nhớ từng khách quen; hệ gợi ý làm lại việc đó bằng dữ liệu, cho hàng triệu người.")

    # ── PHẦN 2
    s = d.divider(2, QUESTIONS[1][0], "Hai loại phản hồi, một ma trận gần như trống, và đuôi dài")

    s = d.slide()
    d.title(s, "Phản hồi tường minh thì hiếm, phản hồi ẩn thì nhiễu")
    cw2 = CW // 2 - Inches(0.15)
    cols = [("Tường minh", "người dùng tự nói thích hay không",
             ["chấm 1–5 sao", "bấm Thích / Không thích", "viết đánh giá"], "rõ nghĩa, nhưng ít người làm"),
            ("Ẩn", "suy ra từ hành vi",
             ["bấm xem, xem bao lâu, xem hết hay bỏ dở", "tìm kiếm, thêm vào danh sách", "với sàn TMĐT: xem, thêm giỏ, mua"],
             "rất nhiều, nhưng không nói rõ thích hay không")]
    for i, (h, sub, items, verdict) in enumerate(cols):
        x = L + i * (cw2 + Inches(0.3))
        d.box(s, x, Inches(1.5), cw2, Inches(2.95), fill=TINT)
        T(s, x + Inches(0.35), Inches(1.7), cw2 - Inches(0.7), Inches(0.5), h, size=22, bold=True, color=NAVY)
        T(s, x + Inches(0.35), Inches(2.2), cw2 - Inches(0.7), Inches(0.4), sub, size=15, color=MUTED)
        T(s, x + Inches(0.35), Inches(2.7), cw2 - Inches(0.7), Inches(1.1), items, size=16, bullet=True, space_after=3)
        T(s, x + Inches(0.35), Inches(3.85), cw2 - Inches(0.7), Inches(0.45), verdict, size=16, bold=True,
          color=RED if i else GREEN)
    T(s, L, Inches(4.65), CW, Inches(0.9),
      [[("Ví dụ: ", {"bold": True, "color": NAVY}),
        ("năm 2017 Netflix bỏ thang 5 sao, chuyển sang Thích / Không thích. Trong thử nghiệm, số lượt đánh giá tăng 200%.", {})]],
      size=16)
    d.question(s, "Câu hỏi 2.", "Vì sao một thang đo “kém chi tiết” hơn lại cho dữ liệu tốt hơn?", y=Inches(5.75), h=Inches(0.85))
    d.source(s, "Netflix (5/4/2017), “Goodbye Stars, Hello Thumbs”")
    d.notes(s, "Đáp: bấm 2 nút dễ hơn nên nhiều người đánh giá hơn; ít nhãn mà đủ người vẫn tốt hơn nhãn chi tiết nhưng thưa. "
               "Hầu hết hệ gợi ý thực tế dựa chủ yếu vào phản hồi ẩn.")

    s = d.slide()
    d.title(s, "Bốn loại dữ liệu đầu vào")
    rows = [["", "Về người dùng", "Về sản phẩm"],
            ["Hành vi", "Phản hồi ẩn: xem, bấm, mua, tìm kiếm", "—"],
            ["Đánh giá", "Phản hồi tường minh: sao, thích, nhận xét", "—"],
            ["Thuộc tính", "Khu vực, thiết bị, ngôn ngữ (có từ khi đăng ký)", "Thể loại, diễn viên, mô tả (có từ khi lên kệ)"]]
    d.table(s, L, Inches(1.5), CW, rows, [Inches(1.9), Inches(5.27), Inches(5.16)], size=16,
            row_h=Inches(0.62), first_col_bold=True)
    T(s, L, Inches(4.25), CW, Inches(0.9),
      [[("Lọc cộng tác ", {"bold": True, "color": NAVY}), ("chỉ cần hai hàng đầu. ", {}),
        ("Lọc theo nội dung ", {"bold": True, "color": NAVY}), ("cần thêm thuộc tính sản phẩm.", {})]], size=18)
    d.question(s, "Thảo luận.", "Một người chưa từng bấm xem một phim. Có thể kết luận người đó không thích phim đó không?",
               y=Inches(5.75), h=Inches(0.85))
    d.notes(s, "Đáp: không. Có thể người đó chưa từng thấy phim. Ô trống trộn lẫn “không thích” và “chưa thấy”, "
               "nên không được coi là 0 điểm. Học liệu mục 8.2.")

    s = d.slide()
    d.title(s, "Ma trận người dùng × sản phẩm gần như trống")
    d.pic(s, "ma_tran", L, Inches(1.45), w=Inches(10.6))
    d.takeaway(s, "Hai người bất kỳ thường chỉ chung vài sản phẩm, nên rất khó tìm “người giống mình”", y=Inches(5.95))
    d.source(s, "MovieLens 1M (GroupLens): 6.040 người × 3.706 phim, 95,5% ô trống")
    d.notes(s, "Mỗi hàng một người, mỗi cột một sản phẩm. MovieLens 1M là dữ liệu công khai cùng loại với Netflix Prize: "
               "1 triệu lượt chấm trên 22,4 triệu ô. Sàn TMĐT còn thưa hơn nhiều vì danh mục lớn hơn.")

    s = d.slide()
    d.title(s, "Vài sản phẩm chiếm phần lớn tương tác: đuôi dài")
    d.pic(s, "duoi_dai", L, Inches(1.4), h=Inches(4.3))
    xr = Inches(7.6); wr = Inches(5.2)
    T(s, xr, Inches(1.6), wr, Inches(0.5), "Phần đầu", size=20, bold=True, color=NAVY)
    T(s, xr, Inches(2.05), wr, Inches(0.8), "20% phim nhận 65% lượt chấm. Ai cũng biết, không cần gợi ý.", size=16)
    T(s, xr, Inches(3.05), wr, Inches(0.5), "Phần đuôi", size=20, bold=True, color=NAVY)
    T(s, xr, Inches(3.5), wr, Inches(0.8), "Rất nhiều phim ít người biết. Đây là nơi hệ gợi ý tạo thêm giá trị.", size=16)
    d.question(s, "Thảo luận.", "Gợi ý cùng 10 phim nổi tiếng nhất cho mọi người: cách đơn giản này tốt hay kém?",
               y=Inches(5.95), h=Inches(0.85))
    d.source(s, "MovieLens 1M")
    d.notes(s, "Giữ câu trả lời của lớp. Phần 4 sẽ thấy: gợi ý phổ biến nhất là đường cơ sở rất khó vượt, chính vì đuôi dài.")

    # ── PHẦN 3
    s = d.divider(3, QUESTIONS[2][0], "Lọc theo nội dung, lọc cộng tác, và vì sao chúng cho kết quả khác nhau")

    s = d.slide()
    d.title(s, "Một trang chủ là nhiều thuật toán ghép lại")
    d.pic(s, "thanh_phan", L, Inches(1.45), w=Inches(11.2))
    d.takeaway(s, "Không có “một thuật toán gợi ý”: mỗi hàng giải một bài toán con khác nhau", y=Inches(6.1))
    d.source(s, "Netflix: 80% giờ xem đến từ gợi ý, 20% từ tìm kiếm (Gomez-Uribe & Hunt, 2015)")
    d.notes(s, "Sáu thành phần theo Gomez-Uribe & Hunt (2015); thành phần thứ sáu (page generation) chọn hàng nào xuất hiện và đặt ở đâu. "
               "Hôm nay tập trung vào hàng “Vì bạn đã xem…”: tìm sản phẩm giống một sản phẩm cho trước.")

    s = d.slide()
    d.title(s, "“Giống phim Ma trận” có thể hiểu theo hai cách")
    for i, (h, steps) in enumerate([
            ("Cách 1. Giống về mô tả", ["An thích Ma trận", "Ma trận: hành động, viễn tưởng, giật gân", "gợi ý phim có cùng đặc điểm"]),
            ("Cách 2. Giống về người xem", ["An thích Ma trận", "những người cũng thích Ma trận", "gợi ý phim họ thích mà An chưa xem"])]):
        y = Inches(1.6) + i * Inches(1.75)
        T(s, L, y, CW, Inches(0.5), h, size=20, bold=True, color=NAVY)
        bw = Inches(3.75)
        for j, st in enumerate(steps):
            x = L + j * (bw + Inches(0.54))
            d.box(s, x, y + Inches(0.55), bw, Inches(0.85), fill=YEL_T if j == 2 else TINT)
            T(s, x + Inches(0.2), y + Inches(0.55), bw - Inches(0.4), Inches(0.85), st, size=16,
              anchor="m", bold=(j == 2), color=NAVY if j == 2 else INK)
            if j < 2:
                T(s, x + bw, y + Inches(0.55), Inches(0.54), Inches(0.85), "›", size=28, color=MUTED,
                  align="c", anchor="m")
    d.question(s, "Câu hỏi 3.", "Mỗi cách cần dữ liệu gì? Cách nào có thể cho kết quả bất ngờ hơn?",
               y=Inches(5.4), h=Inches(0.85))
    d.notes(s, "Để lớp trả lời trước. Cách 1 cần mô tả sản phẩm. Cách 2 cần lịch sử của mọi người, không cần biết phim thuộc thể loại gì, "
               "nên tìm được liên hệ mà mô tả không có, vì vậy bất ngờ hơn. Cách 2 chính là ý tưởng của Hill 1995.")

    s = d.slide()
    d.title(s, "Hai họ phương pháp gợi ý")
    cw2 = CW // 2 - Inches(0.15)
    for i, (h, en, idea, need) in enumerate([
            ("Lọc theo nội dung", "content-based filtering", "Bạn thích một sản phẩm thì sẽ thích sản phẩm giống nó",
             "mô tả sản phẩm + lịch sử của chính bạn"),
            ("Lọc cộng tác", "collaborative filtering", "Người giống bạn thích gì thì bạn cũng có thể thích",
             "lịch sử của nhiều người dùng")]):
        x = L + i * (cw2 + Inches(0.3))
        d.box(s, x, Inches(1.5), cw2, Inches(3.9), fill=NAVY)
        T(s, x + Inches(0.4), Inches(1.75), cw2 - Inches(0.8), Inches(0.55), h, size=26, bold=True, color=WHITE)
        T(s, x + Inches(0.4), Inches(2.3), cw2 - Inches(0.8), Inches(0.4), en, size=15, italic=True, color=YEL)
        T(s, x + Inches(0.4), Inches(3.0), cw2 - Inches(0.8), Inches(1.0), idea, size=22, color=WHITE)
        T(s, x + Inches(0.4), Inches(4.5), cw2 - Inches(0.8), Inches(0.6),
          [[("Cần: ", {"bold": True, "color": YEL}), (need, {})]], size=16, color="DCE4F0")
    T(s, L, Inches(5.8), CW, Inches(0.8),
      [[("Hệ lai ", {"bold": True, "color": NAVY}),
        ("kết hợp cả hai. Đây là cách các nền tảng lớn thực sự vận hành (phần 5).", {})]], size=18)
    d.notes(s, "Học liệu mục 8.2. Tiếp theo: chạy cả hai cách trên một ví dụ nhỏ với dữ liệu thật để thấy chúng “nghĩ” khác nhau thế nào.")

    s = d.slide()
    d.title(s, "Ví dụ chạy tay: nên gợi ý phim nào cho An?")
    d.pic(s, "ma_tran_8x8", L, Inches(1.4), w=Inches(11.0))
    d.question(s, "Đoán trước.", "An chưa xem 3 phim (ô vàng). Phim nào nên lên màn hình của An?", y=Inches(5.95), h=Inches(0.85))
    d.source(s, "Điểm thật trong MovieLens 1M; tên người là tên đặt thay mã số")
    d.notes(s, "An = người dùng số 403. Ô càng đậm điểm càng cao, “?” là chưa chấm. Học liệu Bảng 8.1. "
               "Ta xem cách 1 và cách 2 trả lời câu này thế nào — chú ý ý tưởng, không cần tính.")

    s = d.slide()
    d.title(s, "Lọc theo nội dung: dựng hồ sơ rồi so khớp")
    steps = [("Hồ sơ sản phẩm", "thể loại, diễn viên, đạo diễn, từ khoá mô tả"),
             ("Hồ sơ người dùng", "người đó thích đặc điểm nào, ghét đặc điểm nào"),
             ("So khớp", "sản phẩm khớp hồ sơ nhất được gợi ý")]
    sw = (CW - Inches(0.6)) // 3
    for i, (h, t) in enumerate(steps):
        x = L + i * (sw + Inches(0.3))
        d.box(s, x, Inches(1.5), sw, Inches(2.0), fill=TINT)
        T(s, x + Inches(0.3), Inches(1.65), Inches(0.8), Inches(0.7), str(i + 1), size=32, bold=True, color=YEL)
        T(s, x + Inches(0.3), Inches(2.3), sw - Inches(0.6), Inches(0.45), h, size=20, bold=True, color=NAVY)
        T(s, x + Inches(0.3), Inches(2.75), sw - Inches(0.6), Inches(0.7), t, size=16, color=INK)
    T(s, L, Inches(3.95), CW, Inches(0.45), "Hai chi tiết quyết định chất lượng", size=18, bold=True, color=NAVY)
    T(s, L, Inches(4.45), CW, Inches(1.8), [
        [("Mô tả dạng văn bản: ", {"bold": True}),
         ("giữ những từ xuất hiện nhiều ở sản phẩm này nhưng ít ở sản phẩm khác (TF-IDF).", {})],
        [("So với mức chấm thường ngày của chính người đó: ", {"bold": True}),
         ("người dễ tính chấm 4 sao có thể là chê, người khó tính chấm 3 sao có thể là khen.", {})]],
      size=18, bullet=True, space_after=10)
    d.notes(s, "Học liệu mục 8.4. Khái niệm “chuẩn hoá”: trừ điểm trung bình của chính người đó trước khi xét thích hay ghét. "
               "Độ khớp đo bằng độ tương đồng cosine (công thức 8.1) — trên lớp chỉ cần hiểu: càng gần 1 càng khớp.")

    s = d.slide()
    d.title(s, "Với An, lọc theo nội dung chọn Titanic")
    d.pic(s, "ho_so_an", L, Inches(1.45), w=Inches(11.3))
    d.takeaway(s, "An chấm 5 sao một phim tình cảm, hồ sơ nghiêng về Tình cảm, Titanic khớp nhất", y=Inches(5.6), h=Inches(0.8))
    d.notes(s, "Hồ sơ của An dựng từ 5 phim An đã chấm, so với mức trung bình 3,8 sao của An. An thích Người đàn bà đẹp (tình cảm) "
               "và Ma trận (giật gân), chê Jura (phiêu lưu). Titanic (chính kịch, tình cảm) khớp nhất.")

    s = d.slide()
    d.title(s, "Lọc theo nội dung: được gì, mất gì")
    rows = [["Được", "Mất"],
            ["Chỉ cần mô tả sản phẩm và lịch sử của một người", "Mô tả nghèo thì gợi ý kém: Ma trận và Kẻ hủy diệt 2 cùng thể loại, không phân biệt được"],
            ["Giải thích được: “vì bạn thích phim tình cảm”", "Quá chuyên biệt: chỉ gợi ý thứ giống cái đã thích"],
            ["Gợi ý được sản phẩm mới, chưa ai dùng", "Người dùng mới chưa có hồ sơ"]]
    d.table(s, L, Inches(1.5), CW, rows, [CW // 2, CW - CW // 2], size=19, row_h=Inches(1.15),
            highlight={(3, 0)})
    d.notes(s, "Ô tô vàng: giữ ý “gợi ý được sản phẩm mới” cho phần 5 (khởi động lạnh).")

    s = d.slide()
    d.title(s, "Lọc cộng tác: người từng đồng ý sẽ tiếp tục đồng ý")
    steps = [("Đo độ giống", "hai người có cùng thích, cùng chê những sản phẩm giống nhau không"),
             ("Chọn láng giềng", "k người giống mình nhất"),
             ("Dự đoán", "láng giềng thích gì mà mình chưa dùng")]
    for i, (h, t) in enumerate(steps):
        x = L + i * (sw + Inches(0.3))
        d.box(s, x, Inches(1.5), sw, Inches(2.2), fill=TINT)
        T(s, x + Inches(0.3), Inches(1.65), Inches(0.8), Inches(0.7), str(i + 1), size=32, bold=True, color=YEL)
        T(s, x + Inches(0.3), Inches(2.3), sw - Inches(0.6), Inches(0.45), h, size=20, bold=True, color=NAVY)
        T(s, x + Inches(0.3), Inches(2.75), sw - Inches(0.6), Inches(0.9), t, size=16, color=INK)
    T(s, L, Inches(4.15), CW, Inches(0.45), "Hai biến thể", size=18, bold=True, color=NAVY)
    rows = [["Theo người dùng", "“Những người giống bạn cũng thích…”"],
            ["Theo sản phẩm", "“Người thích sản phẩm này cũng thích…”, chính là hàng “Vì bạn đã xem…”"]]
    d.table(s, L, Inches(4.65), CW, rows, [Inches(2.6), CW - Inches(2.6)], size=18, header=False,
            row_h=Inches(0.75), first_col_bold=True)
    d.notes(s, "Học liệu mục 8.5. Theo sản phẩm ổn định hơn và tính trước được, nên phổ biến hơn trong thực tế (Amazon 2003).")

    s = d.slide()
    d.title(s, "Giống nhau là cùng lên cùng xuống, không phải cùng số sao")
    d.pic(s, "do_giong", L, Inches(1.45), w=Inches(11.0))
    d.takeaway(s, "Đo bằng độ tương đồng có chuẩn hoá: từ −1 (ngược gu) đến 1 (cùng gu)", y=Inches(5.6), h=Inches(0.8))
    d.notes(s, "Khoa chấm cao hơn An ở mọi phim (Khoa dễ tính hơn) nhưng hai đường song song nên rất giống. Minh thích Jura, An chê Jura: gu ngược. "
               "Con số trong ngoặc là độ tương đồng cosine có chuẩn hoá (gần với hệ số tương quan Pearson): từ −1 (ngược) đến 1 (giống). Học liệu Bảng 8.2.")

    s = d.slide()
    d.title(s, "Với An, lọc cộng tác chọn Kẻ hủy diệt 2")
    d.pic(s, "lang_gieng", L, Inches(1.45), w=Inches(11.3))
    d.takeaway(s, "Ba người giống An nhất đều chấm Kẻ hủy diệt 2 rất cao", y=Inches(5.6), h=Inches(0.8))
    d.notes(s, "Láng giềng: Khoa, Hà, Chi. Trung bình của họ: Vua sư tử 3,00; Kẻ hủy diệt 2 4,67; Titanic 3,00. Học liệu công thức 8.3 cho kết quả cùng thứ tự.")

    s = d.slide()
    d.title(s, "Cùng dữ liệu, hai phương pháp, hai câu trả lời")
    for i, (h, film, why, fill, fc) in enumerate([
            ("Lọc theo nội dung", "Titanic", "vì An thích một phim tình cảm", TINT, NAVY),
            ("Lọc cộng tác", "Kẻ hủy diệt 2", "vì người giống An đều chấm rất cao", YEL, NAVY)]):
        x = L + i * (cw2 + Inches(0.3))
        d.box(s, x, Inches(1.6), cw2, Inches(2.6), fill=fill)
        T(s, x, Inches(1.85), cw2, Inches(0.45), h, size=18, color=fc, align="c")
        T(s, x, Inches(2.35), cw2, Inches(0.9), film, size=40, bold=True, color=fc, align="c", anchor="m")
        T(s, x, Inches(3.35), cw2, Inches(0.5), why, size=16, color=INK, align="c")
    d.question(s, "Thảo luận.", "Lọc cộng tác không biết Kẻ hủy diệt 2 thuộc thể loại gì. Nó dựa vào đâu? "
               "Nếu phải chọn một, bạn kiểm chứng phương pháp nào đúng bằng cách nào?", y=Inches(4.6), h=Inches(1.2))
    d.notes(s, "Những người thích Ma trận như An cũng thích Kẻ hủy diệt 2. Kiểm chứng: thử nghiệm A/B (Bài 7) — phần 6 sẽ thấy Netflix làm đúng như vậy.")

    s = d.slide()
    d.title(s, "Lọc cộng tác tự học ra thể loại từ hành vi")
    d.pic(s, "cum_phim", L, Inches(1.4), w=Inches(10.4))
    d.takeaway(s, "Nó tìm lại được thể loại, và thêm cả những liên hệ mà mô tả không có", y=Inches(6.0))
    d.source(s, "Lọc cộng tác theo sản phẩm trên MovieLens 1M")
    d.notes(s, "Phim có ít nhất 300 lượt chấm. Giống Toy Story nhất: Toy Story 2, A Bug’s Life; giống Titanic: Pretty Woman, Ghost; "
               "giống Ma trận: Star Wars, Kẻ hủy diệt 2. Trong 5 phim giống Toy Story nhất có Nhà tù Shawshank (chính kịch). Học liệu Ví dụ 8.4.")

    s = d.slide()
    d.title(s, "Ba lỗi hay gặp khi đo độ giống")
    errs = [("Chỉ đếm vài sản phẩm chung", "Hai người chung 3 phim, chấm trùng nhau, bị coi là “giống hệt”.",
             "Đòi hỏi đủ số sản phẩm chung; không dùng khoảng cách thô."),
            ("Không trừ mức dễ tính", "Điểm sao luôn dương nên ai trông cũng giống ai.",
             "So với mức chấm thường ngày của từng người."),
            ("Xếp ngược chiều", "Độ giống càng lớn càng giống; xếp tăng dần là chọn đúng người khác mình nhất.",
             "In danh sách láng giềng ra và đọc thử.")]
    rows = [["Lỗi", "Chuyện gì xảy ra", "Cách tránh"]] + [list(e) for e in errs]
    d.table(s, L, Inches(1.5), CW, rows, [Inches(3.0), Inches(5.0), CW - Inches(8.0)], size=18,
            row_h=Inches(1.3), first_col_bold=True)
    d.notes(s, "Ba lỗi lấy từ ví dụ An (học liệu mục 8.6): Hà có khoảng cách Manhattan 0 với An chỉ vì 3 phim chung; cosine trên điểm gốc "
               "của mọi người với An đều 0,885–1,000; xếp tăng dần chọn Minh, Dũng, Bình nên gợi ý sai thành Titanic. Lỗi 3 có thật trong một notebook bài giảng phổ biến.")

    s = d.slide()
    d.title(s, "Hai phương pháp bù cho nhau")
    rows = [["", "Lọc theo nội dung", "Lọc cộng tác"],
            ["Cần", "mô tả sản phẩm", "lịch sử của nhiều người"],
            ["Mạnh", "giải thích được; gợi ý được sản phẩm mới", "không cần mô tả; gợi ý bất ngờ"],
            ["Yếu", "mô tả nghèo; quá chuyên biệt", "dữ liệu thưa; tính toán nặng"],
            ["Người dùng mới", "không gợi ý được", "không gợi ý được"],
            ["Sản phẩm mới", "gợi ý được", "không gợi ý được"]]
    d.table(s, L, Inches(1.5), CW, rows, [Inches(2.6), (CW - Inches(2.6)) // 2, CW - Inches(2.6) - (CW - Inches(2.6)) // 2],
            size=19, row_h=Inches(0.85), first_col_bold=True, highlight={(5, 1)})
    d.notes(s, "Ghép hai cột lại là ra ý tưởng hệ lai (phần 5).")

    # ── PHẦN 4
    s = d.divider(4, QUESTIONS[3][0], "Chỉ số kỹ thuật, đường cơ sở, và khoảng cách tới mục tiêu kinh doanh")

    s = d.slide()
    d.title(s, "Netflix Prize: 1 triệu USD để giảm 10% sai số dự đoán")
    d.pic(s, "prize", L, Inches(1.4), w=Inches(11.6))
    d.question(s, "Câu hỏi 4.", "Netflix trả 1 triệu USD rồi không dùng mô hình thắng giải. Cuộc thi thành công hay thất bại?",
               y=Inches(5.95), h=Inches(0.85))
    d.notes(s, "Dữ liệu cuộc thi: hơn 100 triệu lượt chấm của 480.189 thành viên cho 17.770 phim. Sai số đo bằng RMSE so với Cinematch (0,9525 → 0,8567). "
               "Gợi ý trả lời: thành công một phần — có được phân rã ma trận và danh tiếng; thất bại ở chỗ mục tiêu RMSE không gắn với mục tiêu kinh doanh, và rủi ro dữ liệu cá nhân.")

    s = d.slide()
    d.title(s, "Phân rã ma trận: đặt người và sản phẩm lên cùng một bản đồ")
    d.pic(s, "ban_do", L, Inches(1.4), h=Inches(4.5))
    xr = Inches(8.6); wr = Inches(4.23)
    for i, (h, t) in enumerate([("Mỗi sản phẩm một vị trí", "theo vài “yếu tố ẩn” do mô hình tự tìm ra"),
                                ("Mỗi người một vị trí", "theo khẩu vị"),
                                ("Gần nhau là hợp nhau", "gợi ý những phim nằm gần An")]):
        y = Inches(1.6) + i * Inches(1.25)
        T(s, xr, y, wr, Inches(0.45), h, size=18, bold=True, color=NAVY)
        T(s, xr, y + Inches(0.45), wr, Inches(0.6), t, size=16)
    d.source(s, "Koren, Bell & Volinsky (2009). Hình minh hoạ ý tưởng")
    d.notes(s, "Vị trí trong hình không tính từ dữ liệu. Mô hình thật dùng 20–100 yếu tố, không có tên sẵn; người phân tích đọc ra ý nghĩa sau, ví dụ “hành động ↔ tình cảm”. "
               "Ma trận 6.040 × 3.706 được thay bằng hai bảng nhỏ: vị trí của người và vị trí của phim.")

    s = d.slide()
    d.title(s, "Đánh giá offline: học trên quá khứ, chấm trên tương lai")
    y = Inches(1.6); hb = Inches(1.0)
    d.box(s, L, y, int(CW * 0.64), hb, fill=NAVY, radius=0)
    T(s, L + Inches(0.3), y, int(CW * 0.64) - Inches(0.6), hb,
      [[("Dùng để học: ", {"bold": True}), ("80% tương tác sớm nhất của mỗi người", {})]], size=18, color=WHITE, anchor="m")
    d.box(s, L + int(CW * 0.64), y, CW - int(CW * 0.64), hb, fill=YEL, radius=0)
    T(s, L + int(CW * 0.64) + Inches(0.3), y, CW - int(CW * 0.64) - Inches(0.6), hb,
      [[("Dùng để chấm: ", {"bold": True}), ("20% muộn nhất", {})]], size=18, color=NAVY, anchor="m")
    T(s, L, Inches(2.9), Inches(6.0), Inches(0.45), "Đo gì", size=20, bold=True, color=NAVY)
    T(s, L, Inches(3.4), Inches(6.0), Inches(2.2), [
        [("Độ chính xác P@K: ", {"bold": True}), ("trong K sản phẩm gợi ý, bao nhiêu sản phẩm người đó thật sự dùng sau này.", {})],
        [("Độ phủ danh mục: ", {"bold": True}), ("bao nhiêu phần trăm danh mục từng được gợi ý cho ai đó.", {})]],
      size=18, bullet=True, space_after=10)
    T(s, Inches(7.0), Inches(2.9), Inches(5.83), Inches(0.45), "So với gì", size=20, bold=True, color=NAVY)
    T(s, Inches(7.0), Inches(3.4), Inches(5.83), Inches(2.2), [
        [("Gợi ý ngẫu nhiên: ", {"bold": True}), ("mức sàn.", {})],
        [("Gợi ý phổ biến nhất: ", {"bold": True}), ("cùng một danh sách cho mọi người.", {})]],
      size=18, bullet=True, space_after=10)
    d.takeaway(s, "Mô hình không vượt được gợi ý phổ biến nhất thì chưa đáng triển khai", y=Inches(5.75), h=Inches(0.8))
    d.notes(s, "Tách theo thời gian để mô hình không “nhìn trước” tương lai. P@10 = tỷ lệ phim trúng trong 10 phim gợi ý (học liệu công thức 8.4).")

    s = d.slide()
    d.title(s, "Trên MovieLens, gợi ý phổ biến nhất đã khó vượt")
    d.pic(s, "so_sanh", L, Inches(1.4), w=Inches(11.0))
    d.takeaway(s, "Phổ biến nhất gấp 16 lần ngẫu nhiên; lọc cộng tác chỉ hơn phổ biến nhất 24–28%", y=Inches(5.75), h=Inches(0.8))
    d.source(s, "MovieLens 1M, 5.956 người, P@10 và độ phủ danh mục")
    d.notes(s, "Liên hệ câu hỏi ở slide Đuôi dài. Lọc theo nội dung chỉ với 18 thể loại rất yếu — mô tả càng nghèo, cách 1 càng kém. "
               "Độ chính xác và độ phủ đánh đổi nhau. Học liệu Bảng 8.3.")

    s = d.slide()
    d.title(s, "Hai bài học từ Netflix Prize")
    for i, (h, body, link) in enumerate([
            ("Chỉ số kỹ thuật không phải mục tiêu",
             "Mô hình thắng giải không được dùng vì phần chính xác tăng thêm không bù được chi phí kỹ thuật để vận hành nó.",
             "Ý nghĩa thống kê khác ý nghĩa kinh doanh (Bài 7)."),
            ("Dữ liệu “ẩn danh” chưa chắc ẩn danh",
             "Narayanan & Shmatikov (2008) nhận diện lại một số thành viên bằng cách đối chiếu với điểm họ công khai trên IMDb.",
             "Dùng dữ liệu cá nhân đúng mục đích (Nghị định 13/2023).")]):
        x = L + i * (cw2 + Inches(0.3))
        d.box(s, x, Inches(1.5), cw2, Inches(3.6), fill=TINT)
        T(s, x + Inches(0.4), Inches(1.75), cw2 - Inches(0.8), Inches(0.9), h, size=22, bold=True, color=NAVY)
        T(s, x + Inches(0.4), Inches(2.7), cw2 - Inches(0.8), Inches(1.4), body, size=16)
        T(s, x + Inches(0.4), Inches(4.25), cw2 - Inches(0.8), Inches(0.7), link, size=16, bold=True, color=RED)
    d.takeaway(s, "Netflix đánh giá thuật toán bằng thử nghiệm A/B, với chỉ số chính là tỷ lệ giữ chân", y=Inches(5.45), h=Inches(0.8))
    d.source(s, "Netflix Tech Blog (2012); Narayanan & Shmatikov, IEEE S&P 2008")
    d.notes(s, "Ma trận chấm phim thưa đến mức gần như là “dấu vân tay” của mỗi người. Sau vụ kiện và câu hỏi từ FTC, Netflix huỷ cuộc thi thứ hai (3/2010).")

    # ── PHẦN 5
    s = d.divider(5, QUESTIONS[4][0], "Khởi động lạnh, và vì sao hệ thống thực tế luôn là hệ lai")

    s = d.slide()
    d.title(s, "Ngày đầu ra mắt, sản phẩm mới chưa có một lượt tương tác")
    T(s, L, Inches(1.45), CW, Inches(0.6),
      "Netflix thêm phim mới liên tục, kể cả phim tự sản xuất. Sàn TMĐT nhập hàng mới mỗi ngày.", size=18)
    rows = [["Phương pháp", "Hỏi gì", "Với sản phẩm mới"],
            ["Lọc cộng tác", "Ai đã dùng sản phẩm này?", "Chưa có ai, nên không gợi ý được"],
            ["Lọc theo nội dung", "Sản phẩm có đặc điểm gì?", "Có ngay từ mô tả, nên gợi ý được"]]
    d.table(s, L, Inches(2.25), CW, rows, [Inches(3.0), Inches(4.5), CW - Inches(7.5)], size=18,
            row_h=Inches(0.75), first_col_bold=True, highlight={(2, 2)})
    d.question(s, "Câu hỏi 5.", "Còn một người vừa đăng ký, chưa xem gì thì sao? Bạn sẽ hiển thị gì cho họ?",
               y=Inches(5.0), h=Inches(0.85))
    d.notes(s, "Sản phẩm mới: lọc cộng tác không dùng được (chưa ai xem); lọc theo nội dung dùng được ngay. "
               "Người dùng mới: cả hai đều khó, nên dùng gợi ý phổ biến, hoặc hỏi chọn vài sản phẩm yêu thích khi đăng ký. Đây là vấn đề khởi động lạnh.")

    s = d.slide()
    d.title(s, "Khởi động lạnh được giải bằng hệ lai")
    T(s, L, Inches(1.45), CW, Inches(0.5),
      "Thử trên MovieLens: giả lập 300 phim mới. Bao nhiêu phần trăm ô gợi ý dành cho phim mới?", size=18)
    maxw = Inches(6.5)
    d.hbar(s, L, Inches(2.2), "Lọc cộng tác", "≈ 0%", 0.0, maxw / 0.09 * 0.09, "8C9AAD", h=Inches(0.65), vsize=28)
    d.hbar(s, L, Inches(3.05), "Lọc theo nội dung", "9%", 1.0, maxw, NAVY, h=Inches(0.65), vsize=28)
    T(s, L, Inches(4.15), CW, Inches(1.0), [
        [("Hệ lai: ", {"bold": True, "color": NAVY}),
         ("trộn điểm của hai cách, hoặc dùng nội dung cho sản phẩm mới và chuyển sang cộng tác khi đã đủ dữ liệu.", {})]],
      size=18)
    d.takeaway(s, "Trang chủ Netflix là một hệ lai nhiều tầng: mỗi hàng một thuật toán", y=Inches(5.45), h=Inches(0.8))
    d.notes(s, "Số chính xác: lọc cộng tác 0,003% số ô, lọc theo nội dung 9,0% (phim mới chiếm 8,1% danh mục). Học liệu Ví dụ 8.5 và mục 8.9.")

    # ── PHẦN 6
    s = d.divider(6, QUESTIONS[5][0], "Từ độ chính xác đến tiền: giữ chân, doanh thu, lãi")

    s = d.slide()
    d.title(s, "Doanh nghiệp đo hệ gợi ý bằng tác động lên khách hàng")
    rows = [["Netflix công bố", "Con số", "Ý nghĩa"],
            ["Giờ xem đến từ gợi ý", "80%", "phần lớn nhu cầu được tạo ra trên màn hình, không phải từ tìm kiếm"],
            ["Số phim được xem đáng kể", "gấp khoảng 4", "cá nhân hoá kéo người xem ra khỏi nhóm phim đầu bảng"],
            ["Tỷ lệ huỷ thuê bao", "giảm vài điểm %", "chỉ số chính mà mọi thử nghiệm A/B nhắm tới"],
            ["Giá trị ước tính", "> 1 tỷ USD/năm", "chi phí giữ chân mà Netflix tránh được"]]
    d.table(s, L, Inches(1.5), CW, rows, [Inches(3.4), Inches(2.4), CW - Inches(5.8)], size=17,
            row_h=Inches(0.72), first_col_bold=True)
    d.question(s, "Câu hỏi 6.", "Vì sao chỉ số chính là giữ chân chứ không phải độ chính xác? Nếu thuật toán mới làm thời gian xem tăng "
               "nhưng chỉ đẩy vài phim nổi tiếng, cần nhìn thêm chỉ số nào?", y=Inches(5.45), h=Inches(1.15))
    d.source(s, "Gomez-Uribe & Hunt (2015), số liệu do Netflix tự công bố")
    d.notes(s, "“Gấp khoảng 4”: kích thước danh mục hiệu dụng khi có cá nhân hoá gấp khoảng 4 lần khi không có (ở vị trí xếp hạng trung vị). "
               "Đáp: giữ chân gắn trực tiếp với doanh thu thuê bao; độ chính xác chỉ là phương tiện. Cần nhìn thêm độ phủ danh mục — chỉ số bảo vệ như Bài 7.")

    s = d.slide()
    d.title(s, "Hệ gợi ý tự tạo ra dữ liệu huấn luyện của chính nó")
    d.pic(s, "vong", L, Inches(1.5), h=Inches(4.9))
    xr = Inches(7.3); wr = Inches(5.53)
    for i, (h, t) in enumerate([("Vòng phản hồi", "sản phẩm được gợi ý nhiều thì được dùng nhiều, rồi lại được gợi ý nhiều hơn"),
                                ("Phá vòng", "theo dõi độ phủ; dành một số ô để thăm dò sản phẩm ít được thấy"),
                                ("Quyết định bằng A/B test", "chỉ số chính gắn với tiền: giữ chân, doanh thu, lãi"),
                                ("Phù hợp nhất chưa chắc có lợi nhất", "mỗi sản phẩm mang lại mức lãi khác nhau (Bài 9)")]):
        y = Inches(1.5) + i * Inches(1.3)
        T(s, xr, y, wr, Inches(0.4), h, size=19, bold=True, color=NAVY)
        T(s, xr, y + Inches(0.42), wr, Inches(0.8), t, size=16)
    d.notes(s, "Học liệu mục 8.9–8.10. Lãi kỳ vọng của một ô = mức phù hợp × tỷ lệ chuyển từ hiển thị sang mua × lãi gộp mỗi đơn. "
               "Bài 9 thêm ràng buộc kinh doanh vào bài toán chọn K ô.")

    # ── Tổng kết
    s = d.slide("blue")
    d.title(s, "Sáu câu hỏi, sáu câu trả lời", color=WHITE)
    answers = ["Chọn K sản phẩm cho từng người: một quyết định phân bổ",
               "Phản hồi ẩn và tường minh; ma trận rất thưa, có đuôi dài",
               "Nội dung: so khớp hồ sơ. Cộng tác: người giống mình thích gì",
               "So với đường cơ sở; nhìn cả độ chính xác lẫn độ phủ",
               "Sản phẩm mới dùng nội dung; kết hợp hai cách thành hệ lai",
               "Thử nghiệm A/B với chỉ số gắn với tiền"]
    y0, rh = Inches(1.6), Inches(0.82)
    for i, ((q, _), a) in enumerate(zip(QUESTIONS, answers)):
        y = y0 + i * rh
        T(s, L, y, Inches(0.6), rh, str(i + 1), size=24, bold=True, color=YEL, anchor="m")
        T(s, Inches(1.2), y, Inches(5.0), rh, q, size=17, bold=True, color=WHITE, anchor="m")
        T(s, Inches(6.4), y, Inches(6.43), rh, a, size=17, color="DCE4F0", anchor="m")
    d.notes(s, "Che cột phải, cho lớp tự trả lời từng câu trước khi hiện đáp án.")

    s = d.slide()
    d.title(s, "Áp dụng ngoài Netflix")
    qs = ["Một sàn thương mại điện tử vừa nhập 500 mẫu áo mới, chưa ai mua. Dùng phương pháp nào để gợi ý chúng? Vì sao?",
          "Hàng “Khách mua sản phẩm này cũng mua…” trên một sàn TMĐT thuộc phương pháp nào? Nó cần dữ liệu gì?",
          "Mô hình mới chính xác hơn 2% nhưng chỉ gợi ý 10% danh mục thay vì 40%. Bạn có khuyến nghị triển khai ngay không?"]
    for i, q in enumerate(qs):
        y = Inches(1.55) + i * Inches(1.6)
        d.box(s, L, y, CW, Inches(1.35), fill=TINT)
        T(s, L + Inches(0.35), y, Inches(0.7), Inches(1.35), str(i + 1), size=32, bold=True, color=NAVY, anchor="m")
        T(s, L + Inches(1.1), y, CW - Inches(1.5), Inches(1.35), q, size=19, anchor="m")
    d.notes(s, "Làm theo nhóm 5 phút, mỗi nhóm trình bày một câu.")

    if teacher:
        s = d.slide()
        d.title(s, "Đáp án (bản giảng viên)")
        rows = [["Câu", "Đáp án"],
                ["1", "Lọc theo nội dung (kiểu dáng, màu, chất liệu, giá). Lọc cộng tác cần lượt mua nên chưa gợi ý được sản phẩm mới."],
                ["2", "Lọc cộng tác theo sản phẩm. Cần lịch sử mua (hoặc xem) của nhiều khách, không cần mô tả sản phẩm."],
                ["3", "Chưa. Độ phủ giảm mạnh làm hẹp danh mục, dễ rơi vào vòng phản hồi. Kiểm chứng bằng A/B test với chỉ số kinh doanh, độ phủ là chỉ số bảo vệ."]]
        d.table(s, L, Inches(1.5), CW, rows, [Inches(1.0), CW - Inches(1.0)], size=19, row_h=Inches(1.3))

    return d.prs


if __name__ == "__main__":
    build(True).save(f"{OUT}/AA6060_Bai08_Slides.pptx")
    build(False).save(f"{OUT}/AA6060_Bai08_Slides_SinhVien.pptx")
    print("ok")
